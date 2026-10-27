import io
import os
import uuid
from pathlib import Path
from typing import Optional

try:
    import pymupdf as fitz
except ImportError:
    import fitz

from docx import Document as WordDocument
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Document, DocumentPermission, User
from app.schemas import (
    ActionMessageResponse,
    DocumentDetailResponse,
    DocumentPermissionResponse,
    DocumentResponse,
    DocumentUploadResponse,
    GrantPermissionRequest,
)
from app.security import get_current_user

router = APIRouter(prefix="/documents", tags=["Documents"])

# Private local storage. Kept out of version control and not served statically.
STORAGE_DIR = Path(
    os.getenv(
        "DOCUMENT_STORAGE_DIR",
        str(Path(__file__).resolve().parents[2] / "private_uploads"),
    )
)
MAX_FILE_SIZE = int(os.getenv("MAX_UPLOAD_SIZE_BYTES", 10 * 1024 * 1024))  # 10 MB default
ALLOWED_EXTENSIONS = {".pdf", ".docx"}


def extract_pdf_text(file_path: Path) -> list[dict]:
    """Extract text from each PDF page while preserving page numbers."""
    pages = []

    try:
        pdf_bytes = file_path.read_bytes()
        with fitz.open(stream=pdf_bytes, filetype="pdf") as pdf:
            if pdf.is_encrypted:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Password-protected PDFs are not supported.",
                )

            for page_number, page in enumerate(pdf, start=1):
                pages.append({
                    "page": page_number,
                    "text": page.get_text("text").strip(),
                })
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded PDF is invalid or unreadable.",
        )

    return pages


def extract_docx_text(file_path: Path) -> list[dict]:
    """Extract paragraphs from a Word document."""
    try:
        docx_bytes = io.BytesIO(file_path.read_bytes())
        document = WordDocument(docx_bytes)
        text = "\n".join(
            paragraph.text for paragraph in document.paragraphs
            if paragraph.text.strip()
        )

        # DOCX paragraphs do not have reliable page numbers.
        return [{"page": None, "text": text}]
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded Word document is invalid or unreadable.",
        )


@router.post(
    "/upload",
    response_model=DocumentUploadResponse,
    status_code=status.HTTP_201_CREATED,
    summary="Upload and process a company document",
)
def upload_document(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    Secure document upload endpoint.
    - Requires authenticated user via Bearer token.
    - Accepts only PDF (.pdf) and Word (.docx) files.
    - Enforces file size limit while streaming to disk.
    - Safely extracts text without exposing contents publicly.
    - Stores file with random UUID filename in private directory.
    - Cleans up uploaded files and rolls back DB on any failure.
    """
    original_name = Path(file.filename or "").name
    extension = Path(original_name).suffix.lower()

    if not original_name or len(original_name) > 255:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid filename.",
        )

    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Only PDF and DOCX files are supported.",
        )

    STORAGE_DIR.mkdir(parents=True, exist_ok=True)
    stored_name = f"{uuid.uuid4().hex}{extension}"
    stored_path = STORAGE_DIR / stored_name

    size = 0

    try:
        # Stream to disk while enforcing the configured size limit.
        with stored_path.open("xb") as destination:
            while True:
                chunk = file.file.read(1024 * 1024)
                if not chunk:
                    break

                size += len(chunk)
                if size > MAX_FILE_SIZE:
                    raise HTTPException(
                        status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                        detail=f"File exceeds the {MAX_FILE_SIZE // (1024 * 1024)} MB limit.",
                    )

                destination.write(chunk)

        if size == 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The uploaded file is empty.",
            )

        if extension == ".pdf":
            pages = extract_pdf_text(stored_path)
        else:
            pages = extract_docx_text(stored_path)

        extracted_text = "\n\n".join(
            f"[Page {item['page']}]\n{item['text']}"
            if item["page"] is not None else item["text"]
            for item in pages if item["text"]
        )

        if not extracted_text.strip():
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="No text was extracted. Scanned PDFs may require OCR.",
            )

        record = Document(
            title=original_name,
            filename=original_name,
            storage_path=str(stored_path),
            owner_id=current_user.id,
            status="uploaded",
        )

        db.add(record)
        db.commit()
        db.refresh(record)

        return DocumentUploadResponse(
            message="Document uploaded and text extracted successfully.",
            document_id=record.id,
            filename=record.filename,
            size_bytes=size,
            page_count=len(pages) if extension == ".pdf" else None,
            status=record.status,
            extraction_status="completed",
        )

    except HTTPException:
        db.rollback()
        if stored_path.exists():
            stored_path.unlink(missing_ok=True)
        raise
    except Exception:
        db.rollback()
        if stored_path.exists():
            stored_path.unlink(missing_ok=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Document upload failed.",
        )
    finally:
        file.file.close()


@router.get(
    "",
    response_model=list[DocumentResponse],
    summary="List authorized documents",
)
@router.get(
    "/",
    response_model=list[DocumentResponse],
    include_in_schema=False,
)
def list_documents(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    List only documents the current user is authorized to access.
    - Employees: Only documents they own or have explicit read permission for.
    - Administrators: System documents visible for permission management.
      Explicitly returns is_owner and can_read to enforce the policy that an admin
      role alone does not automatically grant read access to confidential documents.
    """
    if current_user.role == "admin":
        docs = db.query(Document).order_by(Document.created_at.desc(), Document.id.desc()).all()
        result = []
        for doc in docs:
            is_owner = (doc.owner_id == current_user.id)
            can_read = is_owner or any(
                p.user_id == current_user.id and p.can_read for p in doc.permissions
            )
            result.append(
                DocumentResponse(
                    id=doc.id,
                    title=doc.title,
                    filename=doc.filename,
                    owner_id=doc.owner_id,
                    status=doc.status,
                    created_at=doc.created_at,
                    is_owner=is_owner,
                    can_read=can_read,
                )
            )
        return result

    permitted_doc_ids = (
        select(DocumentPermission.document_id)
        .filter(
            DocumentPermission.user_id == current_user.id,
            DocumentPermission.can_read.is_(True),
        )
    )

    docs = (
        db.query(Document)
        .filter(
            or_(
                Document.owner_id == current_user.id,
                Document.id.in_(permitted_doc_ids),
            )
        )
        .order_by(Document.created_at.desc(), Document.id.desc())
        .all()
    )

    result = []
    for doc in docs:
        is_owner = (doc.owner_id == current_user.id)
        result.append(
            DocumentResponse(
                id=doc.id,
                title=doc.title,
                filename=doc.filename,
                owner_id=doc.owner_id,
                status=doc.status,
                created_at=doc.created_at,
                is_owner=is_owner,
                can_read=True,
            )
        )
    return result


@router.get(
    "/{document_id}",
    response_model=DocumentDetailResponse,
    summary="Retrieve safe document metadata",
)
def get_document(
    document_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    Retrieve safe metadata after verifying server-side authorization.
    - Prevents IDOR vulnerabilities by strictly checking permissions on the server.
    - Safe metadata only; never exposes storage_path, extracted text, or credentials.
    - Authorized users: Document owners, users with granted read permission, or administrators.
    - Unauthorized employees receive 403 Forbidden.
    - Nonexistent documents receive 404 Not Found.
    """
    if document_id <= 0:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found.",
        )

    doc = db.query(Document).filter(Document.id == document_id).first()
    if not doc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found.",
        )

    is_owner = (doc.owner_id == current_user.id)
    can_manage = is_owner or (current_user.role == "admin")

    has_read_permission = False
    if not is_owner:
        perm = (
            db.query(DocumentPermission)
            .filter(
                DocumentPermission.document_id == doc.id,
                DocumentPermission.user_id == current_user.id,
            )
            .first()
        )
        has_read_permission = bool(perm and perm.can_read)

    if not (is_owner or has_read_permission or can_manage):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="You do not have permission to access this document.",
        )

    return DocumentDetailResponse(
        id=doc.id,
        title=doc.title,
        filename=doc.filename,
        owner_id=doc.owner_id,
        status=doc.status,
        created_at=doc.created_at,
        is_owner=is_owner,
        can_manage_permissions=can_manage,
    )


@router.post(
    "/{document_id}/permissions",
    response_model=DocumentPermissionResponse,
    status_code=status.HTTP_201_CREATED,
    summary="Grant document read permission",
)
def grant_document_permission(
    document_id: int,
    request: GrantPermissionRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    Grant read access, restricted to the document owner or an authorized administrator.
    - Enforces authorization policy: only document owners or admins can grant access.
    - Prevents users from granting themselves access.
    - Rejects invalid user IDs, nonexistent documents, and owner redundant grants.
    - Handles duplicate permissions gracefully by updating existing record.
    - Rolls back cleanly on database failures.
    """
    if document_id <= 0:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found.",
        )

    doc = db.query(Document).filter(Document.id == document_id).first()
    if not doc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found.",
        )

    can_manage = (doc.owner_id == current_user.id) or (current_user.role == "admin")
    if not can_manage:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="You do not have permission to manage permissions for this document.",
        )

    if request.user_id is None and request.email is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Either user_id or email must be provided.",
        )

    target_user = None
    if request.user_id is not None:
        if request.user_id <= 0:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Target user not found.",
            )
        target_user = db.query(User).filter(User.id == request.user_id).first()
    elif request.email is not None:
        target_user = db.query(User).filter(User.email == request.email.strip().lower()).first()

    if not target_user:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Target user not found.",
        )

    # Security: Prevent users from granting themselves access
    if target_user.id == current_user.id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Users cannot grant permissions to themselves.",
        )

    # Security: Document owner already has full access
    if target_user.id == doc.owner_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Document owner already has full access.",
        )

    try:
        perm = (
            db.query(DocumentPermission)
            .filter(
                DocumentPermission.document_id == doc.id,
                DocumentPermission.user_id == target_user.id,
            )
            .first()
        )
        if perm:
            perm.can_read = request.can_read
            perm.granted_by = current_user.id
        else:
            perm = DocumentPermission(
                document_id=doc.id,
                user_id=target_user.id,
                can_read=request.can_read,
                granted_by=current_user.id,
            )
            db.add(perm)

        db.commit()
        db.refresh(perm)

        return DocumentPermissionResponse(
            id=perm.id,
            document_id=perm.document_id,
            user_id=perm.user_id,
            user_email=target_user.email,
            can_read=perm.can_read,
            granted_by=perm.granted_by,
            created_at=perm.created_at,
        )
    except Exception:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Failed to grant permission due to database error.",
        )


@router.delete(
    "/{document_id}/permissions/{user_id}",
    response_model=ActionMessageResponse,
    status_code=status.HTTP_200_OK,
    summary="Revoke document permission",
)
def revoke_document_permission(
    document_id: int,
    user_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """
    Revoke document access permission for a user.
    - Restricted to the document owner or an authorized administrator.
    - Ensures revoked access takes effect immediately.
    - Rejects invalid IDs, nonexistent permissions, and owner revokes.
    - Rolls back cleanly on database failures.
    """
    if document_id <= 0:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found.",
        )

    doc = db.query(Document).filter(Document.id == document_id).first()
    if not doc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Document not found.",
        )

    can_manage = (doc.owner_id == current_user.id) or (current_user.role == "admin")
    if not can_manage:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="You do not have permission to manage permissions for this document.",
        )

    if user_id <= 0:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="User not found.",
        )

    target_user = db.query(User).filter(User.id == user_id).first()
    if not target_user:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="User not found.",
        )

    if target_user.id == doc.owner_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Cannot revoke permissions from the document owner.",
        )

    perm = (
        db.query(DocumentPermission)
        .filter(
            DocumentPermission.document_id == doc.id,
            DocumentPermission.user_id == user_id,
        )
        .first()
    )
    if not perm:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Permission not found for this user and document.",
        )

    try:
        db.delete(perm)
        db.commit()

        return ActionMessageResponse(
            message="Permission revoked successfully.",
            document_id=doc.id,
            user_id=user_id,
        )
    except Exception:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Failed to revoke permission due to database error.",
        )
